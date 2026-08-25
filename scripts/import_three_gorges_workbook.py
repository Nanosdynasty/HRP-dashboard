"""Normalize the reviewed Three Gorges Reservoir Watch workbook for the app.

The source workbook remains outside the repository.  This importer preserves
source gaps, duplicate-date flags and row-level provenance in a compact JSON
dataset that can be shipped with the dashboard.
"""
from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def clean(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return value


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.source, read_only=True, data_only=True)
    daily = workbook["Daily Data"]
    headers = [str(value) for value in next(daily.iter_rows(values_only=True))]
    key_map = {
        "Date": "date",
        "Upstream Level (m)": "upstream_level_m",
        "Downstream Level (m)": "downstream_level_m",
        "Inflow (m3/s)": "inflow_m3s",
        "Outflow (m3/s)": "outflow_m3s",
        "Source Month": "source_month",
        "Source File": "source_file",
        "Source Row Sequence": "source_row_sequence",
        "Quality Flag": "quality_flag",
    }
    observations = []
    for values in daily.iter_rows(min_row=2, values_only=True):
        source = dict(zip(headers, values))
        if not source.get("Date"):
            continue
        observations.append({target: clean(source.get(label)) for label, target in key_map.items()})

    coverage_sheet = workbook["Monthly Coverage"]
    coverage_headers = [str(value) for value in next(coverage_sheet.iter_rows(values_only=True))]
    coverage = []
    for values in coverage_sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(coverage_headers, values))
        if not row.get("Source Month"):
            continue
        coverage.append({
            "source_month": row.get("Source Month"),
            "rows_extracted": row.get("Rows Extracted"),
            "calendar_days": row.get("Calendar Days"),
            "missing_days": row.get("Missing Days"),
            "duplicate_days": row.get("Duplicate Days"),
            "status": row.get("Status"),
            "source_note": row.get("Source Note"),
            "direct_file_url": row.get("Direct File URL"),
        })

    valid_dates = [item["date"] for item in observations]
    payload = {
        "dataset_id": "three-gorges-reservoir-watch-2020-2026",
        "station": "Three Gorges Reservoir",
        "waterbody": "Yangtze River / Three Gorges Reservoir",
        "country": "China",
        "latitude": 30.8231,
        "longitude": 111.0039,
        "source_name": "Probe International – Three Gorges Reservoir Watch",
        "source_page_url": "https://journal.probeinternational.org/reservoir-level-3/",
        "source_provenance": "User-supplied reviewed extraction of the linked monthly source files",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "coverage_start": min(valid_dates),
        "coverage_end": max(valid_dates),
        "known_gaps": [
            "May 2023 has no download link on the source page.",
            "May–July 2026 are linked but contain no extracted rows in the supplied workbook.",
        ],
        "observations": observations,
        "monthly_coverage": coverage,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {len(observations)} observations to {args.output}")


if __name__ == "__main__":
    main()
