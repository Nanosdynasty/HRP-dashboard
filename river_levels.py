"""Official river and reservoir levels for commercially navigable waterways.

The manager deliberately keeps source-specific meaning intact.  Gauge height is
not converted to navigable depth and a hydrological warning is not presented as
a port closure.  PDFs are parsed in memory; only the normalized JSON cache is
stored.
"""
from __future__ import annotations

import asyncio
import csv
import io
import json
import re
import statistics
import unicodedata
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional
from urllib.parse import urljoin

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


NOAA_GAUGES = ("EADM7", "MEMT1", "VCKM6", "BTRL1", "NORL1")
NOAA_GAUGE_URL = "https://api.water.noaa.gov/nwps/v1/gauges/{gauge_id}"
NOAA_HISTORY_URL = "https://api.water.noaa.gov/nwps/v1/gauges/{gauge_id}/stageflow/observed"
NOAA_DOCS_URL = "https://api.water.noaa.gov/nwps/v1/docs/"
PEGEL_URL = (
    "https://pegelonline.wsv.de/webservices/rest-api/v2/stations.json"
    "?waters=RHEIN,DONAU&includeTimeseries=true&includeCurrentMeasurement=true"
    "&includeCharacteristicValues=true"
)
PEGEL_DOCS_URL = "https://pegelonline.wsv.de/webservice/dokuRestapi"
ANA_AMAZON_PAGE = "https://www.gov.br/ana/pt-br/sala-de-situacao/rio-amazonas"
ACP_GATUN_PDF = "https://evtms-rpts.pancanal.com/eng/h2o/GatunWaterIndicators.pdf"
ACP_GATUN_PAGE = "https://evtms-rpts.pancanal.com/eng/h2o/index.html"
INDONESIA_SIHKA_URL = "https://sihka.dev-tunnels.id/api/v1s/infrastructure/asset?id=15"
INDONESIA_SIHKA_PAGE = "https://sihka.sda.pu.go.id/"
THREE_GORGES_DATA_PATH = Path(__file__).parent / "data" / "three_gorges_reservoir_watch.json"
THREE_GORGES_SOURCE_PAGE = "https://journal.probeinternational.org/reservoir-level-3/"

INDONESIA_STATIONS = {
    "AWLR SRIJABO": ("Musi River", "Musi Basin"),
    "MARABAHAN": ("Barito River", "Barito Basin"),
    "POS DUGA AIR SUKALANTING": ("Kapuas River", "Kapuas Basin"),
}

PEGEL_STATIONS = {
    "KAUB", "MANNHEIM", "MAXAU", "DUISBURG-RUHRORT",
    "PASSAU DONAU", "PFELLING",
}

AMAZON_STATIONS = (
    {
        "marker": "MANAUS-RIONEGRO", "station": "Manaus", "waterbody": "Rio Negro",
        "latitude": -3.135, "longitude": -60.025, "country": "Brazil",
    },
    {
        "marker": "OBIDOS-RIOAMAZONAS", "station": "Óbidos", "waterbody": "Amazon River",
        "latitude": -1.916, "longitude": -55.519, "country": "Brazil",
    },
    {
        "marker": "ITACOATIARA-RIOAMAZONAS", "station": "Itacoatiara", "waterbody": "Amazon River",
        "latitude": -3.143, "longitude": -58.444, "country": "Brazil",
    },
    {
        "marker": "SANTAREM-RIOTAPAJOS", "station": "Santarém", "waterbody": "Tapajós River",
        "latitude": -2.438, "longitude": -54.699, "country": "Brazil",
    },
    {
        "marker": "MANACAPURU-RIOSOLIMOES", "station": "Manacapuru", "waterbody": "Solimões River",
        "latitude": -3.299, "longitude": -60.620, "country": "Brazil",
    },
    {
        "marker": "ITAITUBA-RIOTAPAJOS", "station": "Itaituba", "waterbody": "Tapajós River",
        "latitude": -4.276, "longitude": -55.983, "country": "Brazil",
    },
    {
        "marker": "PORTOVELHO-RIOMADEIRA", "station": "Porto Velho", "waterbody": "Madeira River",
        "latitude": -8.760, "longitude": -63.900, "country": "Brazil",
    },
)

# The workspace is deliberately narrower than a general hydrology map.  A row
# is exposed only when its waterway directly supports coastal shipping, ocean
# transit, or regular commercial/bulk inland navigation.  This allowlist also
# prevents a broad upstream feed from silently adding recreational or purely
# flood-monitoring stations later.
COMMERCIAL_WATERWAY_ROLES = {
    "mississippi river": "Lower Mississippi grain, coal and dry-bulk export corridor",
    "ohio river": "Mississippi-system coal, steel and industrial-bulk corridor",
    "missouri river": "Mississippi-system commercial cargo corridor",
    "amazon river": "Amazon basin coastal and agricultural/mineral bulk corridor",
    "rio negro": "Amazon basin commercial navigation and Manaus port corridor",
    "solimões river": "Upper Amazon commercial navigation corridor",
    "madeira river": "Brazilian grain-barge and Amazon export corridor",
    "tapajós river": "Brazilian grain-barge and Amazon export corridor",
    "rhine": "European inland dry-bulk and North Sea port-gateway corridor",
    "danube": "European inland bulk and Black Sea trade corridor",
    "elbe": "Central European inland cargo and North Sea port corridor",
    "moselle": "Rhine-system steel, coal and industrial-bulk corridor",
    "mekong": "Southeast Asian commercial river-transport corridor",
    "musi river": "South Sumatran coal-barge and port corridor",
    "barito river": "Indonesian coal-barge and coastal-transshipment corridor",
    "kapuas river": "Indonesian inland and coastal cargo corridor",
    "st. lawrence river": "Great Lakes–Atlantic ocean-shipping corridor",
    "lake ontario and st. lawrence river": "Great Lakes–Atlantic ocean-shipping corridor",
    "yangtze river": "Major Chinese inland, coastal and dry-bulk shipping corridor",
    "yangtze river / three gorges reservoir": "Yangtze shipping-system trade-critical reservoir",
    "paraná river": "South American grain and dry-bulk export corridor",
    "paraguay river": "Paraguay–Paraná grain, ore and bulk-barge corridor",
    "gatún lake / panama canal": "Ocean-shipping canal reservoir",
}

EXCLUDED_RIVER_COUNTRIES = {"Bangladesh", "India"}


def commercial_trade_role(row: Dict[str, Any]) -> Optional[str]:
    """Return the reviewed commercial role, or None when the row is out of scope."""
    if str(row.get("country") or "").strip() in EXCLUDED_RIVER_COUNTRIES:
        return None
    waterbody = " ".join(str(row.get("waterbody") or "").casefold().split())
    return COMMERCIAL_WATERWAY_ROLES.get(waterbody)


def filter_commercial_waterways(rows: Iterable[Dict[str, Any]]) -> List[Dict[str, Any]]:
    selected: List[Dict[str, Any]] = []
    for source_row in rows:
        role = commercial_trade_role(source_row)
        if not role:
            continue
        row = dict(source_row)
        row["trade_relevance"] = role
        selected.append(row)
    return selected

SOURCE_CATALOG: List[Dict[str, Any]] = [
    {
        "id": "noaa-nwps", "region": "United States", "waterways": "Mississippi and connected US rivers",
        "authority": "NOAA National Water Prediction Service", "access": "Open REST API",
        "status": "connected", "frequency": "Near real time",
        "metrics": "Gauge stage, flow, forecast, flood categories and thresholds", "url": NOAA_DOCS_URL,
    },
    {
        "id": "usace-rivergages", "region": "United States", "waterways": "Mississippi, Ohio, Missouri and USACE waterways",
        "authority": "US Army Corps of Engineers RiverGages", "access": "Public portal / district services",
        "status": "reference", "frequency": "Near real time",
        "metrics": "Stage, flow, locks, dams and water-control observations",
        "url": "https://rivergages.mvr.usace.army.mil/",
    },
    {
        "id": "ana-amazon", "region": "Brazil", "waterways": "Amazon, Negro, Solimões, Madeira and Tapajós",
        "authority": "ANA / SGB National Hydrometeorological Network", "access": "Daily bulletin; API credentials available from ANA",
        "status": "connected", "frequency": "Daily",
        "metrics": "River level, daily change and hydrological reference status", "url": ANA_AMAZON_PAGE,
    },
    {
        "id": "pegelonline", "region": "Germany / Austria", "waterways": "Rhine and upper Danube",
        "authority": "German Federal Waterways and Shipping Administration", "access": "Open REST API",
        "status": "connected", "frequency": "Typically 15 minutes",
        "metrics": "Water level, discharge where available, forecast and gauge datum", "url": PEGEL_DOCS_URL,
    },
    {
        "id": "euris", "region": "Europe", "waterways": "Rhine, Danube, Elbe, Moselle and other major corridors",
        "authority": "European River Information Services (EuRIS)", "access": "Free Open API",
        "status": "connector_ready", "frequency": "Dynamic",
        "metrics": "Water levels, bridge clearance, discharge, fairway and traffic information",
        "url": "https://developer.eurisportal.eu/docs/hydrometeo/hydrometeo-single",
    },
    {
        "id": "mrc-mekong", "region": "Southeast Asia", "waterways": "Mekong",
        "authority": "Mekong River Commission", "access": "Public monitoring portal",
        "status": "connector_ready", "frequency": "Most telemetry stations: 15 minutes",
        "metrics": "Water level and rainfall", "url": "https://portal.mrcmekong.org/monitoring/river-monitoring-telemetry",
    },
    {
        "id": "indonesia-sihka", "region": "Indonesia", "waterways": "Musi, Barito and Kapuas",
        "authority": "Indonesia Ministry of Public Works SIHKA", "access": "Official public telemetry service",
        "status": "connected", "frequency": "Telemetry, typically 5 minutes",
        "metrics": "Gauge water level and official alert thresholds", "url": INDONESIA_SIHKA_PAGE,
    },
    {
        "id": "eccc-wateroffice", "region": "Canada", "waterways": "St. Lawrence and Canadian waterways",
        "authority": "Environment and Climate Change Canada Water Survey", "access": "Public hydrometric downloads/services",
        "status": "connector_ready", "frequency": "Real time / provisional",
        "metrics": "Water level and flow", "url": "https://wateroffice.ec.gc.ca/",
    },
    {
        "id": "ijc-loslr", "region": "Canada / United States", "waterways": "Lake Ontario and St. Lawrence River",
        "authority": "International Joint Commission", "access": "Public weekly regulation summary",
        "status": "connector_ready", "frequency": "Weekly",
        "metrics": "Lake and river levels, outflow and seasonal comparison",
        "url": "https://www.ijc.org/en/loslrb/watershed/regulation-summary",
    },
    {
        "id": "china-mwr-mot", "region": "China", "waterways": "Yangtze and major Chinese rivers/reservoirs",
        "authority": "Ministry of Water Resources / Ministry of Transport", "access": "Public portals and bulletins",
        "status": "portal_only", "frequency": "Intraday / bulletin",
        "metrics": "Major-river levels, reservoir storage, Yangtze water and tide levels",
        "url": "https://hfc.mwr.cn/",
    },
    {
        "id": "three-gorges-watch", "region": "China", "waterways": "Yangtze / Three Gorges Reservoir",
        "authority": "Probe International Reservoir Watch", "access": "User-supplied extracted workbook with row-level source lineage",
        "status": "connected", "frequency": "Daily observations; latest extracted record 30 April 2026",
        "metrics": "Upstream and downstream level, inflow and outflow", "url": THREE_GORGES_SOURCE_PAGE,
    },
    {
        "id": "argentina-ina", "region": "Argentina", "waterways": "Paraná and Paraguay",
        "authority": "Instituto Nacional del Agua", "access": "Public hydrological information portal",
        "status": "connector_ready", "frequency": "Daily / near real time",
        "metrics": "Gauge level, flow and hydrological outlook",
        "url": "https://www.argentina.gob.ar/ina",
    },
    {
        "id": "acp-gatun", "region": "Panama", "waterways": "Gatún Lake / Panama Canal",
        "authority": "Panama Canal Authority", "access": "Official indicator PDF and CSV downloads",
        "status": "connected", "frequency": "Daily",
        "metrics": "Official lake level, permitted transit draft and fresh-water surcharge", "url": ACP_GATUN_PAGE,
    },
]


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _iso_now() -> str:
    return _utc_now().isoformat()


def _safe_number(value: Any) -> Optional[float]:
    try:
        number = float(value)
        return None if number <= -900 else number
    except (TypeError, ValueError):
        return None


def _trend(current: Optional[float], prior: Optional[float]) -> tuple[str, Optional[float]]:
    if current is None or prior is None:
        return "unknown", None
    delta = round(current - prior, 2)
    if abs(delta) < 0.01:
        return "steady", delta
    return ("rising" if delta > 0 else "falling"), delta


def _percentile(values: List[float], fraction: float) -> Optional[float]:
    if not values:
        return None
    ordered = sorted(values)
    position = (len(ordered) - 1) * fraction
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    weight = position - lower
    return round(ordered[lower] * (1 - weight) + ordered[upper] * weight, 2)


def _comparison_fields(
    level: Optional[float], values: Iterable[Any], unit: str,
    basis: str, change_24h: Optional[float] = None,
    change_7d: Optional[float] = None,
) -> Dict[str, Any]:
    clean = [value for value in (_safe_number(item) for item in values) if value is not None]
    if level is None or not clean:
        return {
            "normal_level": None, "normal_low": None, "normal_high": None,
            "normal_unit": unit, "normal_basis": "Not available from this official feed",
            "difference_from_normal": None, "percent_from_normal": None,
            "change_24h": change_24h, "change_7d": change_7d, "change_unit": unit,
            "historical_percentile": None, "recent_min": None, "recent_max": None,
            "recent_window": None, "comparison_status": "unavailable",
        }
    normal = round(statistics.fmean(clean), 2)
    low = _percentile(clean, 0.25)
    high = _percentile(clean, 0.75)
    difference = round(level - normal, 2)
    percent = round(difference / abs(normal) * 100, 1) if normal else None
    rank = round(sum(value <= level for value in clean) / len(clean) * 100, 1)
    comparison = "below_normal" if low is not None and level < low else "above_normal" if high is not None and level > high else "normal"
    return {
        "normal_level": normal, "normal_low": low, "normal_high": high,
        "normal_unit": unit, "normal_basis": basis,
        "difference_from_normal": difference, "percent_from_normal": percent,
        "change_24h": change_24h, "change_7d": change_7d, "change_unit": unit,
        "historical_percentile": rank, "recent_min": round(min(clean), 2),
        "recent_max": round(max(clean), 2), "recent_window": basis,
        "comparison_status": comparison,
    }


def _source_defined_comparison(
    level: Optional[float], unit: str, normal: Optional[float],
    low: Optional[float], high: Optional[float], basis: str,
) -> Dict[str, Any]:
    comparison = "unavailable"
    if level is not None and normal is not None:
        comparison = "below_normal" if low is not None and level < low else "above_normal" if high is not None and level > high else "normal"
    difference = round(level - normal, 2) if level is not None and normal is not None else None
    percent = round(difference / abs(normal) * 100, 1) if difference is not None and normal else None
    return {
        "normal_level": normal, "normal_low": low, "normal_high": high,
        "normal_unit": unit, "normal_basis": basis if normal is not None else "Not available from this official feed",
        "difference_from_normal": difference, "percent_from_normal": percent,
        "change_24h": None, "change_7d": None, "change_unit": unit,
        "historical_percentile": None, "recent_min": None, "recent_max": None,
        "recent_window": None, "comparison_status": comparison,
    }


def _freshness(observed_at: Optional[str]) -> str:
    if not observed_at:
        return "unknown"
    try:
        value = datetime.fromisoformat(observed_at.replace("Z", "+00:00"))
        hours = (_utc_now() - value.astimezone(timezone.utc)).total_seconds() / 3600
        if hours <= 6:
            return "live"
        if hours <= 36:
            return "recent"
        return "stale"
    except (TypeError, ValueError):
        return "unknown"


def _change_from_history(history: List[Dict[str, Any]], current: Optional[float], hours: int) -> Optional[float]:
    if current is None or not history:
        return None
    parsed: List[tuple[datetime, float]] = []
    for item in history:
        value = _safe_number(item.get("primary") if "primary" in item else item.get("value"))
        timestamp = item.get("validTime") or item.get("timestamp")
        try:
            if value is not None and timestamp:
                parsed.append((datetime.fromisoformat(str(timestamp).replace("Z", "+00:00")), value))
        except ValueError:
            continue
    if not parsed:
        return None
    latest_time = max(item[0] for item in parsed)
    target = latest_time - timedelta(hours=hours)
    prior = min(parsed, key=lambda item: abs((item[0] - target).total_seconds()))
    return round(current - prior[1], 2)


def _noaa_record(data: Dict[str, Any], history_payload: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    observed = data.get("status", {}).get("observed") or {}
    forecast = data.get("status", {}).get("forecast") or {}
    flood_category = observed.get("floodCategory") or "not_reported"
    flood = data.get("flood", {}).get("categories") or {}
    level = _safe_number(observed.get("primary"))
    history = list((history_payload or {}).get("data") or [])
    history_values = [_safe_number(item.get("primary")) for item in history]
    comparison = _comparison_fields(
        level, history_values, observed.get("primaryUnit") or "ft",
        "Trailing 30-day observed distribution",
        _change_from_history(history, level, 24), _change_from_history(history, level, 168),
    )
    record = {
        "id": f"noaa-{str(data.get('lid') or '').lower()}",
        "station": data.get("name") or data.get("lid"),
        "waterbody": "Mississippi River",
        "basin": "Mississippi Basin",
        "country": "United States",
        "waterbody_type": "river",
        "latitude": data.get("latitude"), "longitude": data.get("longitude"),
        "observed_at": observed.get("validTime"),
        "level": level, "level_unit": observed.get("primaryUnit") or "ft",
        "discharge": _safe_number(observed.get("secondary")),
        "discharge_unit": observed.get("secondaryUnit"),
        "status": "high" if flood_category not in {"no_flooding", "not_reported"} else "normal",
        "status_label": str(flood_category).replace("_", " ").title(),
        "navigation_status": "Not assessed by this feed",
        "trend": "forecast available" if forecast.get("primary") is not None else "unknown",
        "trend_value": None, "trend_unit": observed.get("primaryUnit") or "ft",
        "forecast_level": _safe_number(forecast.get("primary")),
        "forecast_time": forecast.get("validTime"),
        "navigation_note": "Gauge stage is relative to the local datum; it is not channel depth or vessel draft.",
        "source_name": "NOAA National Water Prediction Service",
        "source_url": f"https://api.water.noaa.gov/gauge/{data.get('lid')}",
        "source_method": "Open REST API",
        "gauge_datum": "Local NOAA gauge datum",
        "quality_note": "The normal range is a trailing 30-day observational comparison, not a long-term seasonal normal.",
        "history": [
            {"observed_at": item.get("validTime"), "level": _safe_number(item.get("primary")), "level_unit": observed.get("primaryUnit") or "ft"}
            for item in history if _safe_number(item.get("primary")) is not None
        ],
        "thresholds": {
            key: value.get("stage") for key, value in flood.items()
            if isinstance(value, dict) and _safe_number(value.get("stage")) is not None
        },
    }
    record.update(comparison)
    record["freshness"] = _freshness(record["observed_at"])
    return record


def _pegel_records(payload: Iterable[Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for station in payload:
        if station.get("shortname") not in PEGEL_STATIONS:
            continue
        water = station.get("water") or {}
        series = next((item for item in station.get("timeseries", []) if item.get("shortname") == "W"), None)
        measurement = (series or {}).get("currentMeasurement") or {}
        if not measurement:
            continue
        state = measurement.get("stateMnwMhw") or measurement.get("stateNswHsw") or "unknown"
        status = "low" if state == "low" else "high" if state == "high" else "normal" if state == "normal" else "unknown"
        observed_at = measurement.get("timestamp")
        characteristic = {item.get("shortname"): item for item in (series or {}).get("characteristicValues", [])}
        mean_item = characteristic.get("MW") or {}
        normal = _safe_number(mean_item.get("value"))
        low = _safe_number((characteristic.get("MNW") or {}).get("value"))
        high = _safe_number((characteristic.get("MHW") or {}).get("value"))
        span = "–".join(filter(None, [mean_item.get("timespanStart"), mean_item.get("timespanEnd")]))
        comparison = _source_defined_comparison(
            _safe_number(measurement.get("value")), (series or {}).get("unit") or "cm",
            normal, low, high, f"Official PEGELONLINE characteristic values{f' ({span})' if span else ''}",
        )
        row = {
            "id": f"pegel-{station.get('uuid')}", "station": station.get("longname") or station.get("shortname"),
            "waterbody": "Rhine" if water.get("shortname") == "RHEIN" else "Danube",
            "basin": "Rhine Basin" if water.get("shortname") == "RHEIN" else "Danube Basin",
            "country": "Germany" if (station.get("agency") or "").upper() != "VIA DONAU" else "Austria",
            "waterbody_type": "river", "latitude": station.get("latitude"), "longitude": station.get("longitude"),
            "river_km": station.get("km"), "observed_at": observed_at,
            "level": _safe_number(measurement.get("value")), "level_unit": (series or {}).get("unit") or "cm",
            "discharge": None, "discharge_unit": None, "status": status,
            "status_label": "Low relative to mean range" if status == "low" else "High relative to mean range" if status == "high" else "Normal range" if status == "normal" else "Threshold not supplied",
            "navigation_status": "Hydrological range only",
            "trend": "unknown", "trend_value": None, "trend_unit": (series or {}).get("unit") or "cm",
            "forecast_level": None, "forecast_time": None,
            "navigation_note": "Use official fairway-depth and notices-to-skippers services before fixing a vessel draft.",
            "source_name": "PEGELONLINE / Federal Waterways and Shipping Administration",
            "source_url": f"https://pegelonline.wsv.de/gast/stammdaten?pegelnr={station.get('number')}",
            "source_method": "Open REST API", "freshness": _freshness(observed_at),
            "gauge_datum": f"{_safe_number(((series or {}).get('gaugeZero') or {}).get('value')) or 'Not supplied'} {((series or {}).get('gaugeZero') or {}).get('unit') or ''}".strip(),
            "quality_note": "Normal and range are official station characteristic values; gauge height is not fairway depth.",
            "extra_metrics": {
                "Equivalent low-water level": f"{(characteristic.get('GlW') or {}).get('value')} {(characteristic.get('GlW') or {}).get('unit')}" if characteristic.get("GlW") else None,
            },
        }
        row.update(comparison)
        rows.append(row)
    return rows


def _ascii_compact(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", "", value).upper()


def _latest_amazon_pdf(page_html: str) -> tuple[str, Optional[str]]:
    matches = re.findall(r'<a[^>]+href=["\']([^"\']+\.pdf)["\'][^>]*>(.*?)</a>', page_html, re.I | re.S)
    candidates: List[tuple[datetime, str, str]] = []
    for href, label_html in matches:
        label = re.sub(r"<[^>]+>", " ", label_html)
        match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", label)
        if match:
            value = datetime(int(match.group(3)), int(match.group(2)), int(match.group(1)), tzinfo=timezone.utc)
            candidates.append((value, urljoin(ANA_AMAZON_PAGE, href), value.date().isoformat()))
    if not candidates:
        raise ValueError("ANA Amazon page did not expose a dated daily bulletin")
    _, url, date = max(candidates, key=lambda item: item[0])
    return url, date


def _amazon_records(pdf_bytes: bytes, source_url: str, bulletin_date: Optional[str]) -> List[Dict[str, Any]]:
    text = "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(pdf_bytes)).pages)
    compact = _ascii_compact(text)
    rows: List[Dict[str, Any]] = []
    for config in AMAZON_STATIONS:
        start = compact.find(config["marker"])
        if start < 0:
            continue
        # Every station block has one ``CODIGO`` heading.  Cutting at the next
        # heading is safer than cutting at the next selected station because
        # the bulletin contains additional non-trading gauges between them.
        current_code = compact.find("CODIGO", start)
        next_code = compact.find("CODIGO", current_code + 6) if current_code >= 0 else -1
        section = compact[start:next_code if next_code > start else len(compact)]
        matches = list(re.finditer(r"(\d{2}/\d{2}/\d{4})(\d+)(NORMAL|ALERTA|ATENCAO|ESTIAGEM)", section))
        if not matches:
            continue
        latest = matches[-1]
        prior = matches[-2] if len(matches) > 1 else None
        level_cm = float(latest.group(2))
        prior_cm = float(prior.group(2)) if prior else None
        trend, delta = _trend(level_cm, prior_cm)
        status_source = latest.group(3)
        status = "low" if status_source == "ESTIAGEM" else "attention" if status_source == "ATENCAO" else "alert" if status_source == "ALERTA" else "normal"
        observed = datetime.strptime(latest.group(1), "%d/%m/%Y").replace(tzinfo=timezone.utc).isoformat()
        history = [
            {
                "observed_at": datetime.strptime(item.group(1), "%d/%m/%Y").replace(tzinfo=timezone.utc).isoformat(),
                "level": round(float(item.group(2)) / 100.0, 2), "level_unit": "m",
            }
            for item in matches
        ]
        current_m = level_cm / 100.0
        comparison = _comparison_fields(
            current_m, [item["level"] for item in history], "m",
            f"Recent observations in official daily bulletin ({len(history)} days)",
            round(delta / 100.0, 2) if delta is not None else None,
            round(current_m - history[-8]["level"], 2) if len(history) >= 8 else None,
        )
        row = {
            "id": f"ana-{re.sub('[^a-z0-9]+', '-', config['station'].lower()).strip('-')}",
            "station": config["station"], "waterbody": config["waterbody"], "basin": "Amazon Basin",
            "country": config["country"], "waterbody_type": "river",
            "latitude": config["latitude"], "longitude": config["longitude"], "observed_at": observed,
            "level": level_cm / 100.0, "level_unit": "m", "discharge": None, "discharge_unit": None,
            "status": status, "status_label": {"ESTIAGEM": "Low-water condition", "ATENCAO": "Attention", "ALERTA": "Alert", "NORMAL": "Normal"}[status_source],
            "navigation_status": "Hydrological bulletin status; navigation impact requires local authority notice",
            "trend": trend, "trend_value": round(delta / 100.0, 2) if delta is not None else None, "trend_unit": "m / day",
            "forecast_level": None, "forecast_time": None,
            "navigation_note": "Daily ANA bulletin level. It is not a guaranteed fairway depth.",
            "source_name": "ANA Amazon Basin Daily Bulletin / SGB network",
            "source_url": source_url, "source_method": "Official bulletin parsed in memory; normalized JSON retained",
            "freshness": _freshness(observed), "bulletin_date": bulletin_date,
            "gauge_datum": "Official ANA station datum",
            "quality_note": "Baseline is the recent observation series printed in this bulletin, not a long-term seasonal normal.",
            "history": history,
        }
        row.update(comparison)
        rows.append(row)
    return rows


def _gatun_record(pdf_bytes: bytes) -> Dict[str, Any]:
    text = "\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(pdf_bytes)).pages)
    level_match = re.search(r"([\d.]+)\s*ft\s*\n\s*Official Gatun Water Level", text, re.I)
    date_match = re.search(r"Official Gatun Water Level\s*\n?for\s*\n?([A-Z][a-z]{2}\.\s*\d{1,2},\s*\d{4})", text, re.I)
    drafts = re.search(r"Neopanamax\s*\nPanamax\s*\n([\d.]+)\s*ft\s+([\d.]+)\s*ft", text, re.I)
    surcharge = re.search(r"([\d.]+)%\s*\n\s*Variable Fresh Water Surcharge", text, re.I)
    if not level_match:
        raise ValueError("ACP Gatún indicator did not expose the official level")
    observed = None
    if date_match:
        observed = datetime.strptime(date_match.group(1).replace(".", ""), "%b %d, %Y").replace(tzinfo=timezone.utc).isoformat()
    level = float(level_match.group(1))
    row = {
        "id": "acp-gatun-lake", "station": "Gatún Lake", "waterbody": "Gatún Lake / Panama Canal",
        "basin": "Panama Canal Watershed", "country": "Panama", "waterbody_type": "reservoir",
        "latitude": 9.18, "longitude": -79.85, "observed_at": observed,
        "level": level, "level_unit": "ft", "discharge": None, "discharge_unit": None,
        "status": "normal", "status_label": "Official operating indicator",
        "navigation_status": "Official maximum transit drafts published",
        "trend": "unknown", "trend_value": None, "trend_unit": "ft",
        "forecast_level": None, "forecast_time": None,
        "navigation_note": "Transit drafts remain subject to Panama Canal Advisories to Shipping.",
        "source_name": "Panama Canal Authority – Gatún Water Level Indicators",
        "source_url": ACP_GATUN_PAGE, "source_method": "Official indicator PDF parsed in memory; normalized JSON retained",
        "freshness": _freshness(observed),
        "extra_metrics": {
            "Neopanamax max draft": f"{drafts.group(1)} ft" if drafts else None,
            "Panamax max draft": f"{drafts.group(2)} ft" if drafts else None,
            "Fresh-water surcharge": f"{surcharge.group(1)}%" if surcharge else None,
        },
    }
    row.update(_comparison_fields(level, [], "ft", ""))
    row["gauge_datum"] = "Panama Canal operating lake elevation datum"
    row["quality_note"] = "The public indicator does not publish a normal-level series in the same file."
    return row


def _indonesia_records(payload: Dict[str, Any]) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for feature in payload.get("features", []):
        properties = feature.get("properties") or {}
        name = str(properties.get("name") or "").upper()
        if name not in INDONESIA_STATIONS:
            continue
        telemetry = properties.get("last_telemetry") or {}
        observations = sorted((str(key), _safe_number(value)) for key, value in telemetry.items())
        observations = [item for item in observations if item[1] is not None]
        if not observations:
            continue
        observed_at, level = observations[-1]
        if not re.search(r"(?:Z|[+-]\d{2}:?\d{2})$", observed_at):
            observed_at = f"{observed_at}+07:00"
        prior = observations[-2][1] if len(observations) > 1 else None
        trend, delta = _trend(level, prior)
        coordinates = (feature.get("geometry") or {}).get("coordinates") or [None, None]
        waterbody, basin = INDONESIA_STATIONS[name]
        alert_data = properties.get("properties") or {}
        thresholds = {
            f"Alert level {number}": f"{value} m"
            for number in range(1, 5)
            if (value := alert_data.get(f"status_siaga_{number}")) not in (None, "")
        }
        row = {
            "id": f"sihka-{properties.get('id')}", "station": properties.get("name"),
            "waterbody": waterbody, "basin": basin, "country": "Indonesia",
            "waterbody_type": "river", "latitude": coordinates[1], "longitude": coordinates[0],
            "observed_at": observed_at, "level": level, "level_unit": "m",
            "status": "normal", "status_label": "Official gauge observation",
            "trend": trend, "trend_value": delta, "trend_unit": "m / 5 min",
            "navigation_note": "SIHKA gauge level uses its local station datum and is not navigable depth.",
            "source_name": "Indonesia Ministry of Public Works SIHKA",
            "source_url": INDONESIA_SIHKA_PAGE, "source_method": "Official public telemetry service",
            "freshness": _freshness(observed_at), "gauge_datum": "Local SIHKA station datum",
            "quality_note": "The public asset feed supplies the latest telemetry but not a historical normal series.",
            "extra_metrics": thresholds,
            "history": [
                {"observed_at": timestamp, "level": value, "level_unit": "m"}
                for timestamp, value in observations
            ],
        }
        row.update(_comparison_fields(level, [], "m", "", change_24h=None, change_7d=None))
        rows.append(row)
    return rows


def _three_gorges_record(path: Path = THREE_GORGES_DATA_PATH) -> Dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    observations = [
        item for item in payload.get("observations", [])
        if item.get("date") and _safe_number(item.get("upstream_level_m")) is not None
    ]
    if not observations:
        raise ValueError("Three Gorges dataset contains no valid upstream reservoir levels")
    observations.sort(key=lambda item: (item["date"], item.get("source_row_sequence") or 0))
    latest = observations[-1]
    latest_date = datetime.fromisoformat(latest["date"])
    level = _safe_number(latest.get("upstream_level_m"))
    seasonal = [
        _safe_number(item.get("upstream_level_m"))
        for item in observations
        if datetime.fromisoformat(item["date"]).month == latest_date.month
        and datetime.fromisoformat(item["date"]).year < latest_date.year
    ]
    seasonal = [value for value in seasonal if value is not None]

    by_date: Dict[str, float] = {}
    for item in observations:
        by_date[item["date"]] = float(item["upstream_level_m"])
    prior_day = by_date.get((latest_date - timedelta(days=1)).date().isoformat())
    prior_week = by_date.get((latest_date - timedelta(days=7)).date().isoformat())
    comparison = _comparison_fields(
        level, seasonal, "m",
        f"Historical {latest_date.strftime('%B')} distribution, 2020–{latest_date.year - 1}",
        round(level - prior_day, 2) if level is not None and prior_day is not None else None,
        round(level - prior_week, 2) if level is not None and prior_week is not None else None,
    )
    recent = [float(item["upstream_level_m"]) for item in observations[-30:]]
    comparison["recent_min"] = round(min(recent), 2)
    comparison["recent_max"] = round(max(recent), 2)
    comparison["recent_window"] = "Latest 30 extracted daily observations"
    trend, trend_value = _trend(level, prior_day)
    downstream = _safe_number(latest.get("downstream_level_m"))
    inflow = _safe_number(latest.get("inflow_m3s"))
    outflow = _safe_number(latest.get("outflow_m3s"))
    net_flow = round(inflow - outflow, 0) if inflow is not None and outflow is not None else None
    history = [
        {
            "observed_at": f"{item['date']}T00:00:00+08:00",
            "level": _safe_number(item.get("upstream_level_m")), "level_unit": "m",
            "downstream_level_m": _safe_number(item.get("downstream_level_m")),
            "inflow_m3s": _safe_number(item.get("inflow_m3s")),
            "outflow_m3s": _safe_number(item.get("outflow_m3s")),
            "source_month": item.get("source_month"), "source_file": item.get("source_file"),
            "quality_flag": item.get("quality_flag"),
        }
        for item in observations
    ]
    row = {
        "id": "china-three-gorges-reservoir", "station": payload.get("station") or "Three Gorges Reservoir",
        "waterbody": payload.get("waterbody") or "Yangtze River / Three Gorges Reservoir",
        "basin": "Yangtze River Basin", "country": "China", "waterbody_type": "reservoir",
        "latitude": payload.get("latitude"), "longitude": payload.get("longitude"),
        "observed_at": f"{latest['date']}T00:00:00+08:00", "date_precision": "day", "level": level, "level_unit": "m",
        "downstream_level": downstream, "downstream_level_unit": "m",
        "inflow": inflow, "outflow": outflow, "flow_unit": "m³/s", "net_flow": net_flow,
        "status": "normal", "status_label": "Historical reservoir observation",
        "trend": trend, "trend_value": trend_value, "trend_unit": "m / day",
        "navigation_note": "Reservoir surface elevation is not channel depth, least available depth or permissible vessel draft.",
        "source_name": payload.get("source_name") or "Probe International – Three Gorges Reservoir Watch",
        "source_url": payload.get("source_page_url") or THREE_GORGES_SOURCE_PAGE,
        "source_method": payload.get("source_provenance") or "User-supplied workbook normalized locally",
        "freshness": _freshness(f"{latest['date']}T00:00:00+08:00"),
        "gauge_datum": "Reservoir surface elevation in metres, as stated by the source workbook",
        "quality_note": (
            "Historical extracted series, not a live official Chinese telemetry feed. Latest usable record is "
            f"{latest['date']}; May 2023 and May–July 2026 are documented source gaps."
        ),
        "history": history,
        "extra_metrics": {
            "Downstream level at Yichang": f"{downstream:.2f} m" if downstream is not None else None,
            "Reservoir inflow": f"{inflow:,.0f} m³/s" if inflow is not None else None,
            "Reservoir outflow": f"{outflow:,.0f} m³/s" if outflow is not None else None,
            "Net inflow less outflow": f"{net_flow:+,.0f} m³/s" if net_flow is not None else None,
            "Extracted coverage": f"{payload.get('coverage_start')} to {payload.get('coverage_end')}",
            "Extracted daily rows": f"{len(payload.get('observations', [])):,}",
        },
    }
    row.update(comparison)
    return row


def _three_gorges_yichang_record(path: Path = THREE_GORGES_DATA_PATH) -> Dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    observations = [
        item for item in payload.get("observations", [])
        if item.get("date") and _safe_number(item.get("downstream_level_m")) is not None
    ]
    if not observations:
        raise ValueError("Three Gorges dataset contains no valid Yichang downstream levels")
    observations.sort(key=lambda item: (item["date"], item.get("source_row_sequence") or 0))
    latest = observations[-1]
    latest_date = datetime.fromisoformat(latest["date"])
    level = _safe_number(latest.get("downstream_level_m"))
    seasonal = [
        _safe_number(item.get("downstream_level_m"))
        for item in observations
        if datetime.fromisoformat(item["date"]).month == latest_date.month
        and datetime.fromisoformat(item["date"]).year < latest_date.year
    ]
    seasonal = [value for value in seasonal if value is not None]
    by_date: Dict[str, float] = {}
    for item in observations:
        by_date[item["date"]] = float(item["downstream_level_m"])
    prior_day = by_date.get((latest_date - timedelta(days=1)).date().isoformat())
    prior_week = by_date.get((latest_date - timedelta(days=7)).date().isoformat())
    comparison = _comparison_fields(
        level, seasonal, "m",
        f"Historical {latest_date.strftime('%B')} distribution, 2020–{latest_date.year - 1}",
        round(level - prior_day, 2) if level is not None and prior_day is not None else None,
        round(level - prior_week, 2) if level is not None and prior_week is not None else None,
    )
    recent = [float(item["downstream_level_m"]) for item in observations[-30:]]
    comparison["recent_min"] = round(min(recent), 2)
    comparison["recent_max"] = round(max(recent), 2)
    comparison["recent_window"] = "Latest 30 extracted daily observations"
    trend, trend_value = _trend(level, prior_day)
    upstream = _safe_number(latest.get("upstream_level_m"))
    inflow = _safe_number(latest.get("inflow_m3s"))
    outflow = _safe_number(latest.get("outflow_m3s"))
    history = [
        {
            "observed_at": f"{item['date']}T00:00:00+08:00",
            "level": _safe_number(item.get("downstream_level_m")), "level_unit": "m",
            "upstream_level_m": _safe_number(item.get("upstream_level_m")),
            "inflow_m3s": _safe_number(item.get("inflow_m3s")),
            "outflow_m3s": _safe_number(item.get("outflow_m3s")),
            "source_month": item.get("source_month"), "source_file": item.get("source_file"),
            "quality_flag": item.get("quality_flag"),
        }
        for item in observations
    ]
    row = {
        "id": "china-yangtze-yichang", "station": "Yangtze River at Yichang",
        "waterbody": "Yangtze River", "basin": "Yangtze River Basin", "country": "China",
        "waterbody_type": "river", "latitude": 30.692, "longitude": 111.286,
        "observed_at": f"{latest['date']}T00:00:00+08:00", "date_precision": "day", "level": level, "level_unit": "m",
        "upstream_level": upstream, "upstream_level_unit": "m",
        "inflow": inflow, "outflow": outflow, "flow_unit": "m³/s",
        "status": "normal", "status_label": "Historical river observation",
        "trend": trend, "trend_value": trend_value, "trend_unit": "m / day",
        "navigation_note": "Yichang water level is not channel depth, least available depth or permissible vessel draft.",
        "source_name": payload.get("source_name") or "Probe International – Three Gorges Reservoir Watch",
        "source_url": payload.get("source_page_url") or THREE_GORGES_SOURCE_PAGE,
        "source_method": payload.get("source_provenance") or "User-supplied workbook normalized locally",
        "freshness": _freshness(f"{latest['date']}T00:00:00+08:00"),
        "gauge_datum": "Yichang downstream level in metres, as stated by the source workbook",
        "quality_note": (
            "Historical extracted series, not a live official Chinese telemetry feed. The source states that this downstream "
            f"level is at Yichang, about 40 km below the dam. Latest usable record is {latest['date']}."
        ),
        "history": history,
        "extra_metrics": {
            "Upstream reservoir level": f"{upstream:.2f} m" if upstream is not None else None,
            "Reservoir inflow": f"{inflow:,.0f} m³/s" if inflow is not None else None,
            "Reservoir outflow": f"{outflow:,.0f} m³/s" if outflow is not None else None,
            "Extracted coverage": f"{payload.get('coverage_start')} to {payload.get('coverage_end')}",
        },
    }
    row.update(comparison)
    return row


class RiverLevelManager:
    def __init__(self, cache_path: Path, refresh_seconds: int = 3600):
        self.cache_path = cache_path
        self.refresh_seconds = refresh_seconds
        self.payload: Dict[str, Any] = {}
        self.last_error: Optional[str] = None
        self.task: Optional[asyncio.Task] = None
        self._lock = asyncio.Lock()
        self._load_cache()

    def _load_cache(self) -> None:
        try:
            data = json.loads(self.cache_path.read_text(encoding="utf-8"))
            if isinstance(data, dict):
                data["rows"] = filter_commercial_waterways(data.get("rows") or [])
                data["sources"] = SOURCE_CATALOG
                data["source_count"] = len(SOURCE_CATALOG)
                data["connected_source_count"] = sum(
                    item["status"] == "connected" for item in SOURCE_CATALOG
                )
                self.payload = data
        except (FileNotFoundError, json.JSONDecodeError, OSError):
            self.payload = {}

    def _save(self) -> None:
        self.cache_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.cache_path.with_suffix(".tmp")
        temporary.write_text(json.dumps(self.payload, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(self.cache_path)

    async def _fetch_noaa(self, client: httpx.AsyncClient) -> List[Dict[str, Any]]:
        responses = await asyncio.gather(*(
            asyncio.gather(
                client.get(NOAA_GAUGE_URL.format(gauge_id=value)),
                client.get(NOAA_HISTORY_URL.format(gauge_id=value)),
            ) for value in NOAA_GAUGES
        ))
        rows = []
        for gauge_response, history_response in responses:
            gauge_response.raise_for_status()
            history_response.raise_for_status()
            rows.append(_noaa_record(gauge_response.json(), history_response.json()))
        return rows

    async def _fetch_pegel(self, client: httpx.AsyncClient) -> List[Dict[str, Any]]:
        response = await client.get(PEGEL_URL)
        response.raise_for_status()
        return _pegel_records(response.json())

    async def _fetch_amazon(self, client: httpx.AsyncClient) -> List[Dict[str, Any]]:
        page = await client.get(ANA_AMAZON_PAGE)
        page.raise_for_status()
        pdf_url, bulletin_date = _latest_amazon_pdf(page.text)
        pdf = await client.get(pdf_url, timeout=90)
        pdf.raise_for_status()
        return await asyncio.to_thread(_amazon_records, pdf.content, pdf_url, bulletin_date)

    async def _fetch_gatun(self, client: httpx.AsyncClient) -> List[Dict[str, Any]]:
        response = await client.get(ACP_GATUN_PDF, timeout=60)
        response.raise_for_status()
        return [await asyncio.to_thread(_gatun_record, response.content)]

    async def _fetch_indonesia(self, client: httpx.AsyncClient) -> List[Dict[str, Any]]:
        response = await client.get(INDONESIA_SIHKA_URL, timeout=60)
        response.raise_for_status()
        return _indonesia_records(response.json())

    async def refresh(self, force: bool = False) -> Dict[str, Any]:
        async with self._lock:
            if not force and self.payload.get("fetched_at"):
                try:
                    fetched = datetime.fromisoformat(self.payload["fetched_at"].replace("Z", "+00:00"))
                    if (_utc_now() - fetched).total_seconds() < self.refresh_seconds:
                        return self.payload
                except (TypeError, ValueError):
                    pass
            errors: List[str] = []
            rows: List[Dict[str, Any]] = []
            headers = {"User-Agent": "HRP-dashboard/4.1 river-level-monitor"}
            async with httpx.AsyncClient(headers=headers, follow_redirects=True, timeout=40) as client:
                results = await asyncio.gather(
                    self._fetch_noaa(client), self._fetch_pegel(client), self._fetch_amazon(client), self._fetch_gatun(client),
                    self._fetch_indonesia(client),
                    asyncio.to_thread(lambda: [_three_gorges_record(), _three_gorges_yichang_record()]),
                    return_exceptions=True,
                )
            names = ("NOAA", "PEGELONLINE", "ANA Amazon", "Panama Canal Authority", "Indonesia SIHKA", "Three Gorges Reservoir Watch")
            for name, result in zip(names, results):
                if isinstance(result, Exception):
                    errors.append(f"{name}: {result}")
                else:
                    rows.extend(result)
            if not rows and self.payload.get("rows"):
                self.last_error = "; ".join(errors) or "No official source returned data"
                return self.payload
            rows = filter_commercial_waterways(rows)
            rows.sort(key=lambda row: (row.get("waterbody") or "", row.get("station") or ""))
            self.payload = {
                "fetched_at": _iso_now(), "rows": rows, "sources": SOURCE_CATALOG,
                "errors": errors, "source_count": len(SOURCE_CATALOG),
                "connected_source_count": sum(item["status"] == "connected" for item in SOURCE_CATALOG),
                "disclaimer": "Gauge height is not navigable depth. Confirm fairway depths, notices to skippers and local restrictions before operational use.",
            }
            self.last_error = "; ".join(errors) if errors else None
            self._save()
            return self.payload

    def start(self) -> None:
        if self.task and not self.task.done():
            return
        self.task = asyncio.create_task(self._run())

    async def _run(self) -> None:
        while True:
            try:
                await self.refresh()
            except asyncio.CancelledError:
                raise
            except Exception as exc:  # pragma: no cover - defensive scheduler guard
                self.last_error = str(exc)
            await asyncio.sleep(self.refresh_seconds)

    async def stop(self) -> None:
        if self.task and not self.task.done():
            self.task.cancel()
            await asyncio.gather(self.task, return_exceptions=True)


def export_rows_csv(rows: Iterable[Dict[str, Any]]) -> str:
    fields = [
        "id", "station", "waterbody", "basin", "country", "waterbody_type", "trade_relevance",
        "latitude", "longitude", "river_km", "observed_at", "level", "level_unit",
        "normal_level", "normal_low", "normal_high", "normal_unit", "normal_basis",
        "difference_from_normal", "percent_from_normal", "comparison_status",
        "change_24h", "change_7d", "change_unit", "historical_percentile",
        "recent_min", "recent_max", "recent_window", "trend", "trend_value", "trend_unit",
        "upstream_level", "upstream_level_unit", "downstream_level", "downstream_level_unit",
        "inflow", "outflow", "net_flow", "flow_unit",
        "status", "status_label", "gauge_datum", "quality_note", "navigation_note",
        "source_name", "source_url", "source_method", "extra_metrics",
    ]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for item in rows:
        row = dict(item)
        row["extra_metrics"] = json.dumps(row.get("extra_metrics") or {}, ensure_ascii=False)
        writer.writerow(row)
    return buffer.getvalue()


def export_river_levels_xlsx(payload: Dict[str, Any], rows: Optional[List[Dict[str, Any]]] = None) -> bytes:
    selected = list(rows if rows is not None else payload.get("rows", []))
    workbook = Workbook()
    current = workbook.active
    current.title = "Current levels"
    fields = [
        "station", "waterbody", "basin", "country", "waterbody_type", "trade_relevance", "observed_at",
        "level", "level_unit", "normal_level", "normal_low", "normal_high", "normal_unit",
        "difference_from_normal", "percent_from_normal", "comparison_status", "normal_basis",
        "change_24h", "change_7d", "change_unit", "historical_percentile",
        "recent_min", "recent_max", "recent_window", "trend", "status_label",
        "upstream_level", "upstream_level_unit", "downstream_level", "downstream_level_unit",
        "inflow", "outflow", "net_flow", "flow_unit",
        "gauge_datum", "quality_note", "latitude", "longitude", "source_name", "source_url",
    ]
    labels = [value.replace("_", " ").title() for value in fields]
    current.append(labels)
    for row in selected:
        current.append([row.get(field) for field in fields])

    history_sheet = workbook.create_sheet("Observation history")
    history_sheet.append([
        "Station", "Waterway", "Country", "Observed at", "Gauge level", "Unit",
        "Related upstream level", "Related downstream level", "Inflow m3/s", "Outflow m3/s",
        "Source month", "Source file", "Quality flag", "Period",
    ])
    for row in selected:
        for item in row.get("history") or []:
            history_sheet.append([
                row.get("station"), row.get("waterbody"), row.get("country"), item.get("observed_at"),
                item.get("level"), item.get("level_unit") or row.get("level_unit"), item.get("upstream_level_m"),
                item.get("downstream_level_m"), item.get("inflow_m3s"), item.get("outflow_m3s"), item.get("source_month"), item.get("source_file"),
                item.get("quality_flag"), item.get("period"),
            ])

    baseline_sheet = workbook.create_sheet("Normal baselines")
    baseline_sheet.append(["Station", "Waterway", "Country", "Normal", "Low", "High", "Unit", "Basis", "Coverage note"])
    for row in selected:
        baseline_sheet.append([
            row.get("station"), row.get("waterbody"), row.get("country"), row.get("normal_level"),
            row.get("normal_low"), row.get("normal_high"), row.get("normal_unit"), row.get("normal_basis"),
            row.get("quality_note"),
        ])

    reference_sheet = workbook.create_sheet("Commercial scope")
    reference_sheet.append(["Country", "Waterway", "Commercial trade role", "Official source"])
    seen_scope = set()
    for row in selected:
        scope_key = (row.get("country"), row.get("waterbody"), row.get("trade_relevance"))
        if scope_key in seen_scope:
            continue
        seen_scope.add(scope_key)
        reference_sheet.append([
            row.get("country"), row.get("waterbody"), row.get("trade_relevance"), row.get("source_name"),
        ])

    source_sheet = workbook.create_sheet("Sources")
    source_fields = ["region", "waterways", "authority", "status", "frequency", "metrics", "access", "url"]
    source_sheet.append([value.title() for value in source_fields])
    for source in payload.get("sources", SOURCE_CATALOG):
        source_sheet.append([source.get(field) for field in source_fields])

    dictionary = workbook.create_sheet("Data dictionary")
    dictionary.append(["Field", "Definition"])
    definitions = {
        "Level": "Latest official gauge or reservoir level in the source's own datum.",
        "Normal level": "Mean of the explicitly stated comparison series; see Normal basis.",
        "Normal low / high": "Interquartile range for observational baselines or official mean-low/mean-high characteristics.",
        "Difference from normal": "Current level minus normal level in the displayed unit.",
        "Percent from normal": "Difference divided by the absolute normal level; omitted where no valid baseline exists.",
        "Historical percentile": "Share of observations in the stated recent window at or below the current level.",
        "Change 24h / 7d": "Current level minus the nearest official observation at the stated lag.",
        "Comparison status": "Below normal, normal, above normal, or unavailable from the official feed.",
        "Gauge datum": "Vertical reference used by the gauge. It is not channel depth.",
        "Trade relevance": "Reviewed reason the waterway is included: coastal shipping, ocean transit, or regular commercial/bulk navigation.",
    }
    for key, value in definitions.items():
        dictionary.append([key, value])

    header_fill = PatternFill("solid", fgColor="173A5E")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True)
        for column in range(1, sheet.max_column + 1):
            values = [str(sheet.cell(row=row, column=column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
            sheet.column_dimensions[get_column_letter(column)].width = min(max(max(map(len, values), default=10) + 2, 11), 42)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
